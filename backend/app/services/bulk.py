import csv
import io
from typing import Iterable

from openpyxl import load_workbook


EXPECTED_COLUMNS = {"company_name", "rc", "nif", "website", "wilaya", "legal_form"}


def _clean_row(row: dict) -> dict[str, str]:
    cleaned: dict[str, str] = {}
    for key, value in row.items():
        if key is None or value is None:
            continue
        normalized_key = str(key).strip().lower()
        if normalized_key in EXPECTED_COLUMNS:
            cleaned[normalized_key] = str(value).strip()
    return cleaned


def parse_bulk_file(data: bytes, *, max_rows: int = 1000) -> list[dict[str, str]]:
    if data.startswith(b"PK\x03\x04"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip().lower() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return []
        rows = (
            _clean_row(dict(zip(headers, values, strict=False)))
            for values in iterator
        )
    else:
        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = (_clean_row(row) for row in reader)

    output: list[dict[str, str]] = []
    for row in rows:
        if not row:
            continue
        output.append(row)
        if len(output) > max_rows:
            raise ValueError("BULK_ROW_LIMIT_EXCEEDED")
    return output
